"""
GT Table Visualizations
"""

from typing import List, Optional
import numpy as np
import pandas as pd
import tempfile
import os
import json
import zipfile
import streamlit as st
from .base import BaseVisualizer

from pathlib import Path

import plotly.express as px
import plotly.graph_objects as go

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, Spacer, PageBreak
from datetime import datetime

from reportlab.platypus import Table, TableStyle, SimpleDocTemplate
from reportlab.lib.pagesizes import letter

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter


class GTTableVisualizer(BaseVisualizer):
    def __init__(self):
        super().__init__()

    
    # Required Columns
    
    REQUIRED_COLUMNS = [
        "Date",
        "Ticker",
        "Model",
        "Prediction",
        "GroundTruth",
        "Confidence",
        "TrustScore",
    ]

    
    # Validate GT Table
    
    def validate_gt_table(
        self,
        df: pd.DataFrame,
    ):
        self.validate_columns(
            df,
            self.REQUIRED_COLUMNS,
        )
        return True
    
    # Search
    
    def search(
        self,
        df,
        keyword,
    ):
        if keyword is None or keyword == "":
            return df
        mask = np.column_stack([
            df[col]
            .astype(str)
            .str.contains(
                keyword,
                case=False,
                na=False,
            )
            for col in df.columns
        ])
        return df.loc[mask.any(axis=1)]
    
    # Filter Model
    
    def filter_model(
        self,
        df,
        model,
    ):
        if model == "All":
            return df
        return df[
            df["Model"] == model
        ]

    
    # Filter Ticker
    
    def filter_ticker(
        self,
        df,
        ticker,
    ):
        if ticker == "All":
            return df
        return df[
            df["Ticker"] == ticker
        ]

    
    # Filter Trust
    
    def filter_trust(
        self,
        df,
        minimum=0.0,
        maximum=1.0,
    ):
        return df[
            (df["TrustScore"] >= minimum)
            &
            (df["TrustScore"] <= maximum)
        ]

    
    # Filter Date
    
    def filter_date(
        self,
        df,
        start_date,
        end_date,
    ):

        return df[
            (df["Date"] >= start_date)
            &
            (df["Date"] <= end_date)
        ]
    
    # Highlight Trust
    
    @staticmethod
    def highlight_trust(value):
        if value >= 0.80:
            return "background-color:#D1FAE5"
        elif value >= 0.60:
            return "background-color:#FEF3C7"
        else:
            return "background-color:#FECACA"

    
    # Highlight Prediction

    @staticmethod
    def highlight_prediction(row):
        if row["Prediction"] == row["GroundTruth"]:
            return [
                "background-color:#DCFCE7"
            ] * len(row)
        else:
            return [
                "background-color:#FEE2E2"
            ] * len(row)

    
    # Summary
    
    def summary(
        self,
        df,
    ):
        return {
            "Rows":
                len(df),

            "Unique Models":
                df["Model"].nunique(),

            "Tickers":
                df["Ticker"].nunique(),

            "Average Trust":
                round(
                    df["TrustScore"].mean(),
                    3,
                ),

            "Average Confidence":
                round(
                    df["Confidence"].mean(),
                    3,
                ),
        }
    
    # Display Summary

    def show_summary(
        self,
        df,
    ):
        summary = self.summary(df)

        c1, c2, c3, c4, c5 = st.columns(5)

        c1.metric(
            "Rows",
            summary["Rows"],
        )

        c2.metric(
            "Models",
            summary["Unique Models"],
        )

        c3.metric(
            "Assets",
            summary["Tickers"],
        )

        c4.metric(
            "Avg Trust",
            summary["Average Trust"],
        )

        c5.metric(
            "Avg Confidence",
            summary["Average Confidence"],
        )

    
    # Interactive GT Table
    
    def show_table(
        self,
        df,
        height=650,
    ):

        styled = (
            df.style
            .apply(
                self.highlight_prediction,
                axis=1,
            )

            .map(
                self.highlight_trust,
                subset=["TrustScore"],
            )
        )

        st.dataframe(
            styled,
            use_container_width=True,
            height=height,
        )
    
    # Download
    
    def download_button(
        self,
        df,
        filename="gt_table.csv",
    ):
        st.download_button(
            label="Download GT Table",
            data=df.to_csv(index=False),
            file_name=filename,
            mime="text/csv",
        )
    
    # Sidebar Filters

    def sidebar_filters(
        self,
        df,
    ):

        st.sidebar.header("GT Table Filters")

        models = [
            "All"
        ] + sorted(
            df["Model"].unique()
        )

        tickers = [
            "All"
        ] + sorted(
            df["Ticker"].unique()
        )

        model = st.sidebar.selectbox(
            "Model",
            models,
        )

        ticker = st.sidebar.selectbox(
            "Ticker",
            tickers,
        )

        trust = st.sidebar.slider(

            "Minimum Trust",
            0.0,
            1.0,
            0.50,
            0.01,
        )

        keyword = st.sidebar.text_input(
            "Search",
            "",
        )

        data = self.filter_model(
            df,
            model,
        )

        data = self.filter_ticker(
            data,
            ticker,
        )

        data = self.filter_trust(
            data,
            trust,
            1.0,
        )

        data = self.search(
            data,
            keyword,
        )
        return data

    
    # Render
    
    def render(
        self,
        df,
    ):
        self.validate_gt_table(df)
        filtered = self.sidebar_filters(df)
        self.show_summary(filtered)
        st.divider()
        self.show_table(filtered)
        self.download_button(filtered)

# Prediction Accuracy

def prediction_accuracy(
    self,
    df,
):
    correct = (
        df["Prediction"] ==
        df["GroundTruth"]
    ).sum()

    total = len(df)

    return round(
        correct / total,
        4,
    )

# Agreement Column

def add_agreement_column(
    self,
    df,
):

    data = df.copy()
    data["Agreement"] = np.where(
        data["Prediction"] ==
        data["GroundTruth"],
        "Correct",
        "Incorrect",
    )
    return data

# Accuracy by Model

def accuracy_by_model(
    self,
    df,
):
    data = self.add_agreement_column(df)
    result = (
        data
        .groupby("Model")["Agreement"]
        .apply(
            lambda x:
            (x == "Correct").mean()
        )
        .reset_index()
    )
    result.columns = [
        "Model",
        "Accuracy",
    ]
    return result


# Accuracy by Ticker

def accuracy_by_ticker(
    self,
    df,
):
    data = self.add_agreement_column(df)
    result = (
        data
        .groupby("Ticker")["Agreement"]
        .apply(
            lambda x:
            (x == "Correct").mean()
        )
        .reset_index()
    )
    result.columns = [
        "Ticker",
        "Accuracy",
    ]
    return result


# Confusion Matrix

def confusion_matrix(
    self,
    df,
):
    return pd.crosstab(
        df["GroundTruth"],
        df["Prediction"],
        margins=True,
    )


# GT vs Prediction Table

def comparison_table(
    self,
    df,
):
    data = self.add_agreement_column(df)
    return data[
        [
            "Date",
            "Ticker",
            "Model",
            "GroundTruth",
            "Prediction",
            "Agreement",
            "Confidence",
            "TrustScore",
        ]
    ]


# Incorrect Predictions

def incorrect_predictions(
    self,
    df,
):
    return df[
        df["Prediction"]
        !=
        df["GroundTruth"]
    ]


# Correct Predictions

def correct_predictions(
    self,
    df,
):
    return df[
        df["Prediction"]
        ==
        df["GroundTruth"]
    ]


# High Confidence Errors

def high_confidence_errors(
    self,
    df,
    threshold=0.90,
):
    return df[
        (
            df["Prediction"]
            !=
            df["GroundTruth"]
        )
        &
        (
            df["Confidence"]
            >= threshold
        )
    ]


# Low Trust Correct Predictions

def low_trust_correct(
    self,
    df,
    threshold=0.50,
):
    return df[
        (
            df["Prediction"]
            ==
            df["GroundTruth"]
        )
        &
        (
            df["TrustScore"]
            < threshold
        )
    ]


# Agreement Summary

def agreement_summary(
    self,
    df,
):
    correct = (
        df["Prediction"]
        ==
        df["GroundTruth"]
    ).sum()
    incorrect = len(df) - correct
    return {
        "Correct":
            int(correct),
        "Incorrect":
            int(incorrect),
        "Accuracy":
            round(
                correct /
                len(df),
                4,
            ),
    }


# Trust vs Agreement

def trust_accuracy_summary(
    self,
    df,
):
    data = self.add_agreement_column(df)
    return (
        data
        .groupby("Agreement")
        [
            [
                "TrustScore",
                "Confidence",
            ]
        ]
        .mean()
        .round(3)
    )


# Prediction Distribution

def prediction_distribution(
    self,
    df,
):
    return (
        df
        .groupby(
            "Prediction"
        )
        .size()
        .reset_index(
            name="Count"
        )
    )


# Ground Truth Distribution

def groundtruth_distribution(
    self,
    df,
):
    return (
        df
        .groupby(
            "GroundTruth"
        )
        .size()
        .reset_index(
            name="Count"
        )
    )


# Model Summary

def model_summary(
    self,
    df,
):
    result = (
        df
        .groupby("Model")
        .agg(
            Predictions=("Prediction", "count"),
            AvgConfidence=("Confidence", "mean"),
            AvgTrust=("TrustScore", "mean"),
        )
        .round(3)
    )
    return result.reset_index()


# Ticker Summary

def ticker_summary(
    self,
    df,
):
    result = (
        df
        .groupby("Ticker")
        .agg(
            Predictions=("Prediction", "count"),
            AvgConfidence=("Confidence", "mean"),
            AvgTrust=("TrustScore", "mean"),
        )
        .round(3)
    )
    return result.reset_index()


# Show Comparison Table

def show_comparison(
    self,
    df,
):
    st.subheader("Ground Truth vs Model Prediction")
    st.dataframe(
        self.comparison_table(df),
        use_container_width=True,
    )


# Show Confusion Matrix

def show_confusion_matrix(
    self,
    df,
):
    st.subheader("Confusion Matrix")
    st.dataframe(
        self.confusion_matrix(df),
        use_container_width=True,
    )

# Show Errors

def show_errors(
    self,
    df,
):
    st.subheader("Prediction Errors")
    st.dataframe(
        self.incorrect_predictions(df),
        use_container_width=True,
    )


# Model Leaderboard

def model_leaderboard(self, df):
    data = self.add_agreement_column(df)
    leaderboard = (
        data
        .groupby("Model")
        .agg(
            Accuracy=("Agreement",
                      lambda x: (x == "Correct").mean()),
            Trust=("TrustScore", "mean"),
            Confidence=("Confidence", "mean"),
            Trades=("Prediction", "count"),
        )
        .round(4)
    )

    leaderboard["Score"] = (
        0.50 * leaderboard["Accuracy"]
        + 0.30 * leaderboard["Trust"]
        + 0.20 * leaderboard["Confidence"]
    )
    leaderboard = (
        leaderboard
        .sort_values(
            "Score",
            ascending=False,
        )
        .reset_index()
    )
    leaderboard.insert(
        0,
        "Rank",
        np.arange(
            1,
            len(leaderboard) + 1,
        ),
    )
    return leaderboard


# Asset Leaderboard

def asset_leaderboard(self, df):
    data = self.add_agreement_column(df)
    leaderboard = (
        data
        .groupby("Ticker")
        .agg(
            Accuracy=("Agreement",
                      lambda x: (x == "Correct").mean()),
            Trust=("TrustScore", "mean"),
            Confidence=("Confidence", "mean"),
            Predictions=("Prediction", "count"),
        )
        .round(4)
        .sort_values(
            "Accuracy",
            ascending=False,
        )
    )
    return leaderboard.reset_index()


# Trade Audit Table

def trade_audit_table(self, df):
    audit = df.copy()
    audit["Correct"] = (
        audit["Prediction"]
        ==
        audit["GroundTruth"]
    )
    audit["Error"] = (
        audit["Prediction"]
        !=
        audit["GroundTruth"]
    )
    audit["RiskFlag"] = np.where(
        (
            audit["Confidence"] > 0.90
        )
        &
        (
            audit["Error"]
        ),
        "HIGH",
        "NORMAL",
    )
    return audit


# GT Scorecard

def gt_scorecard(self, df):
    accuracy = self.prediction_accuracy(df)
    trust = df["TrustScore"].mean()
    confidence = df["Confidence"].mean()
    hallucination = (
        df["Prediction"]
        !=
        df["GroundTruth"]
    ).mean()
    return {
        "Accuracy":
            round(accuracy, 4),
        "Trust":
            round(trust, 4),
        "Confidence":
            round(confidence, 4),
        "Hallucination Rate":
            round(hallucination, 4),
    }


# Model Ranking

def rank_models(self, df):
    board = self.model_leaderboard(df)
    return board.sort_values(
        "Score",
        ascending=False,
    )

# Trust Ranking

def trust_ranking(self, df):
    return (
        df
        .groupby("Model")["TrustScore"]
        .mean()
        .sort_values(
            ascending=False,
        )
        .reset_index()
    )


# Confidence Ranking

def confidence_ranking(self, df):
    return (
        df
        .groupby("Model")["Confidence"]
        .mean()
        .sort_values(
            ascending=False,
        )
        .reset_index()
    )


# Hallucination Ranking

def hallucination_ranking(self, df):
    data = self.add_agreement_column(df)
    ranking = (
        data
        .groupby("Model")["Agreement"]
        .apply(
            lambda x:
            (x == "Incorrect").mean()
        )
        .reset_index()
    )
    ranking.columns = [
        "Model",
        "HallucinationRate",
    ]
    return ranking.sort_values(
        "HallucinationRate"
    )

# Calibration Ranking

def calibration_ranking(self, df):
    ranking = (
        df
        .groupby("Model")
        .apply(
            lambda x:
            np.mean(
                np.abs(
                    x["Confidence"]
                    -
                    x["TrustScore"]
                )
            )
        )
        .reset_index()
    )
    ranking.columns = [
        "Model",
        "CalibrationError",
    ]
    return ranking.sort_values(
        "CalibrationError"
    )

# Overall Audit Report

def audit_report(self, df):
    return {
        "Scorecard":
            self.gt_scorecard(df),
        "Leaderboard":
            self.model_leaderboard(df),
        "Assets":
            self.asset_leaderboard(df),
        "Hallucination":
            self.hallucination_ranking(df),
        "Calibration":
            self.calibration_ranking(df),
    }

# Executive Summary

def executive_summary(self, df):
    score = self.gt_scorecard(df)
    return pd.DataFrame(
        {
            "Metric": [
                "Accuracy",
                "Trust",
                "Confidence",
                "Hallucination",
            ],
            "Value": [
                score["Accuracy"],
                score["Trust"],
                score["Confidence"],
                score["Hallucination Rate"],
            ],
        }
    )

# Export Report

def export_audit_report(
    self,
    df,
    filename="audit_report.csv",
):
    report = self.executive_summary(df)
    st.download_button(
        "Download Audit Report",
        report.to_csv(index=False),
        filename,
        mime="text/csv",
    )

# Show Leaderboard

def show_leaderboard(self, df):
    st.subheader("Model Leaderboard")
    st.dataframe(
        self.model_leaderboard(df),
        use_container_width=True,
    )

# Show Asset Leaderboard

def show_assets(self, df):
    st.subheader("Asset Leaderboard")
    st.dataframe(
        self.asset_leaderboard(df),
        use_container_width=True,
    )

# Show Executive Summary

def show_summary_report(self, df):
    st.subheader("Executive Summary")
    st.dataframe(
        self.executive_summary(df),
        use_container_width=True,
    )

# Full Audit Dashboard

def render_audit_dashboard(self, df):
    self.show_summary(df)
    st.divider()
    self.show_leaderboard(df)
    st.divider()
    self.show_assets(df)
    st.divider()
    self.show_summary_report(df)
    st.divider()
    self.export_audit_report(df)


# Rolling Trust Score

def rolling_trust(
    self,
    df,
    window=30,
):
    data = df.copy()
    data = data.sort_values("Date")
    data["RollingTrust"] = (
        data["TrustScore"]
        .rolling(window)
        .mean()
    )
    return data


# Rolling Accuracy

def rolling_accuracy(
    self,
    df,
    window=30,
):
    data = self.add_agreement_column(df)
    data["Correct"] = (
        data["Agreement"]
        ==
        "Correct"
    ).astype(int)
    data = data.sort_values("Date")
    data["RollingAccuracy"] = (
        data["Correct"]
        .rolling(window)
        .mean()
    )
    return data


# Daily Audit Summary

def daily_summary(
    self,
    df,
):
    data = self.add_agreement_column(df)
    summary = (
        data
        .groupby("Date")
        .agg(
            Accuracy=(
                "Agreement",
                lambda x:
                (x == "Correct").mean(),
            ),
            Trust=(
                "TrustScore",
                "mean",
            ),
            Confidence=(
                "Confidence",
                "mean",
            ),
            Trades=(
                "Prediction",
                "count",
            ),
        )
        .round(4)
    )
    return summary.reset_index()


# Monthly Audit Summary

def monthly_summary(
    self,
    df,
):
    data = df.copy()
    data["Date"] = pd.to_datetime(
        data["Date"]
    )
    data["Month"] = (
        data["Date"]
        .dt.to_period("M")
        .astype(str)
    )
    data = self.add_agreement_column(data)
    summary = (
        data
        .groupby("Month")
        .agg(
            Accuracy=(
                "Agreement",
                lambda x:
                (x == "Correct").mean(),
            ),
            Trust=(
                "TrustScore",
                "mean",
            ),
            Confidence=(
                "Confidence",
                "mean",
            ),
            Trades=(
                "Prediction",
                "count",
            ),
        )
        .round(4)
    )
    return summary.reset_index()


# Trust Heatmap

def trust_heatmap(
    self,
    df,
):
    table = pd.pivot_table(
        df,
        index="Model",
        columns="Ticker",
        values="TrustScore",
        aggfunc=np.mean,
    )
    return table.round(3)

# Accuracy Heatmap

def accuracy_heatmap(
    self,
    df,
):
    data = self.add_agreement_column(df)
    data["Correct"] = (
        data["Agreement"]
        ==
        "Correct"
    ).astype(int)

    table = pd.pivot_table(
        data,
        index="Model",
        columns="Ticker",
        values="Correct",
        aggfunc=np.mean,
    )
    return table.round(3)


# Model Drift

def model_drift(
    self,
    df,
):
    data = df.copy()
    data["Date"] = pd.to_datetime(
        data["Date"]
    )
    data["Month"] = (
        data["Date"]
        .dt.to_period("M")
        .astype(str)
    )
    drift = (
        data
        .groupby(
            [
                "Month",
                "Model",
            ]
        )
        ["TrustScore"]
        .mean()
        .reset_index()
    )
    return drift


# Drift Score

def drift_score(
    self,
    df,
):
    drift = self.model_drift(df)
    result = (
        drift
        .groupby("Model")
        ["TrustScore"]
        .std()
        .reset_index()
    )
    result.columns = [
        "Model",
        "DriftScore",
    ]
    return result.sort_values(
        "DriftScore",
        ascending=False,
    )


# Stable Models

def stable_models(
    self,
    df,
):
    drift = self.drift_score(df)
    return drift.sort_values(
        "DriftScore",
        ascending=True,
    )


# Trust Change

def trust_change(
    self,
    df,
):
    monthly = self.monthly_summary(df)
    monthly["TrustChange"] = (
        monthly["Trust"]
        .diff()
    )
    return monthly


# Accuracy Change

def accuracy_change(
    self,
    df,
):
    monthly = self.monthly_summary(df)
    monthly["AccuracyChange"] = (
        monthly["Accuracy"]
        .diff()
    )
    return monthly


# Risk Flags

def temporal_risk_flags(
    self,
    df,
):
    monthly = self.monthly_summary(df)
    monthly["Risk"] = np.where(
        (
            monthly["Trust"] < 0.60
        )
        |
        (
            monthly["Accuracy"] < 0.55
        ),
        "HIGH",
        "NORMAL",
    )
    return monthly


# Audit Timeline

def audit_timeline(
    self,
    df,
):
    daily = self.daily_summary(df)
    return daily[
        [
            "Date",
            "Accuracy",
            "Trust",
            "Confidence",
            "Trades",
        ]
    ]


# Time Window Comparison

def compare_periods(
    self,
    df,
    start1,
    end1,
    start2,
    end2,
):
    p1 = df[
        (
            df["Date"] >= start1
        )
        &
        (
            df["Date"] <= end1
        )
    ]
    p2 = df[
        (
            df["Date"] >= start2
        )
        &
        (
            df["Date"] <= end2
        )
    ]
    return {
        "Period1":
            self.gt_scorecard(p1),
        "Period2":
            self.gt_scorecard(p2),
    }


# Trust Trend by Model

def trust_trend_by_model(
    self,
    df,
):
    data = df.copy()
    data["Date"] = pd.to_datetime(
        data["Date"]
    )
    data["Month"] = (
        data["Date"]
        .dt.to_period("M")
        .astype(str)
    )
    return (
        data
        .groupby(
            [
                "Month",
                "Model",
            ]
        )
        ["TrustScore"]
        .mean()
        .reset_index()
    )


# Prediction Flow Table

def prediction_flow(
    self,
    df,
):
    """
    Creates a flow table betweenGround Truth -> Prediction.
    """
    flow = (
        df
        .groupby(
            [
                "GroundTruth",
                "Prediction",
            ]
        )
        .size()
        .reset_index(name="Count")
        .sort_values(
            "Count",
            ascending=False,
        )
    )
    return flow


# Model Prediction Counts

def model_prediction_counts(
    self,
    df,
):
    counts = (
        df
        .groupby(
            [
                "Model",
                "Prediction",
            ]
        )
        .size()
        .reset_index(name="Count")
    )
    return counts


# Ground Truth Counts

def groundtruth_counts(
    self,
    df,
):
    counts = (
        df
        .groupby("GroundTruth")
        .size()
        .reset_index(name="Count")
    )
    return counts



# Prediction Transition Matrix

def prediction_transition_matrix(
    self,
    df,
):
    matrix = pd.crosstab(
        df["GroundTruth"],
        df["Prediction"],
    )
    return matrix


# Model Transition Matrix

def model_transition_matrix(
    self,
    df,
):
    table = pd.pivot_table(
        df,
        index="Model",
        columns="Prediction",
        values="Confidence",
        aggfunc="count",
        fill_value=0,
    )
    return table


# Majority Vote

def majority_vote(
    self,
    df,
):
    vote = (
        df
        .groupby(
            [
                "Date",
                "Ticker",
            ]
        )
        ["Prediction"]
        .agg(
            lambda x:
            x.mode().iloc[0]
            if len(x.mode()) > 0
            else np.nan
        )
        .reset_index()
    )
    vote.columns = [
        "Date",
        "Ticker",
        "MajorityPrediction",
    ]
    return vote


# Consensus Score

def consensus_score(
    self,
    df,
):
    def score(group):
        freq = (
            group["Prediction"]
            .value_counts(normalize=True)
            .max()
        )
        return freq
    result = (
        df
        .groupby(
            [
                "Date",
                "Ticker",
            ]
        )
        .apply(score)
        .reset_index(name="Consensus")
    )
    return result


# Consensus with GT

def consensus_vs_groundtruth(
    self,
    df,
):
    consensus = self.majority_vote(df)
    gt = (
        df
        [
            [
                "Date",
                "Ticker",
                "GroundTruth",
            ]
        ]
        .drop_duplicates()
    )
    merged = (
        consensus
        .merge(
            gt,
            on=[
                "Date",
                "Ticker",
            ],
            how="left",
        )
    )
    merged["Correct"] = (
        merged["MajorityPrediction"]
        ==
        merged["GroundTruth"]
    )
    return merged


# Consensus Accuracy

def consensus_accuracy(
    self,
    df,
):
    result = self.consensus_vs_groundtruth(df)
    return round(
        result["Correct"].mean(),
        4,
    )

# Consensus by Model Count

def consensus_distribution(
    self,
    df,
):
    score = self.consensus_score(df)
    score["Level"] = pd.cut(
        score["Consensus"],
        bins=[
            0,
            0.50,
            0.75,
            0.90,
            1.00,
        ],
        labels=[
            "Low",
            "Medium",
            "High",
            "Perfect",
        ],
        include_lowest=True,
    )
    return (
        score
        .groupby("Level")
        .size()
        .reset_index(name="Cases")
    )


# Asset Prediction Flow

def asset_prediction_flow(
    self,
    df,
):
    return (
        df
        .groupby(
            [
                "Ticker",
                "Prediction",
            ]
        )
        .size()
        .reset_index(name="Count")
    )


# Model Prediction Flow

def model_prediction_flow(
    self,
    df,
):
    return (
        df
        .groupby(
            [
                "Model",
                "Prediction",
            ]
        )
        .size()
        .reset_index(name="Count")
    )


# Model Agreement Matrix

def model_agreement_matrix(
    self,
    df,
):
    models = sorted(df["Model"].unique())

    matrix = pd.DataFrame(
        index=models,
        columns=models,
        dtype=float,
    )
    for m1 in models:
        d1 = (
            df[df["Model"] == m1]
            [["Date", "Ticker", "Prediction"]]
            .rename(
                columns={
                    "Prediction": m1
                }
            )
        )
        for m2 in models:
            d2 = (
                df[df["Model"] == m2]
                [["Date", "Ticker", "Prediction"]]
                .rename(
                    columns={
                        "Prediction": m2
                    }
                )
            )
            merged = d1.merge(
                d2,
                on=[
                    "Date",
                    "Ticker",
                ],
               how="inner",
            )
            if len(merged) == 0:
                matrix.loc[m1, m2] = np.nan
            else:
                matrix.loc[m1, m2] = (
                    merged[m1] == merged[m2]
                ).mean()
    return matrix.round(3)


# Model Disagreement Matrix

def model_disagreement_matrix(
    self,
    df,
):
    agreement = self.model_agreement_matrix(df)
    return (1 - agreement).round(3)


# Agreement Graph Edge List

def agreement_edges(
    self,
    df,
    threshold=0.70,
):
    agreement = self.model_agreement_matrix(df)
    edges = []
    models = agreement.index.tolist()
    for i in range(len(models)):
        for j in range(i + 1, len(models)):
            value = agreement.iloc[i, j]
            if pd.notna(value) and value >= threshold:
                edges.append(
                    {
                        "Source": models[i],
                        "Target": models[j],
                        "Weight": value,
                    }
                )
    return pd.DataFrame(edges)


# Agreement Degree

def agreement_degree(
    self,
    df,
    threshold=0.70,
):
    edges = self.agreement_edges(
        df,
        threshold,
    )
    models = sorted(df["Model"].unique())
    degree = []
    for model in models:
        value = (
            (edges["Source"] == model).sum()
            +
            (edges["Target"] == model).sum()
        )
        degree.append(value)
    return pd.DataFrame(
        {
            "Model": models,
            "Degree": degree,
        }
    )


# Consensus Network

def consensus_network(
    self,
    df,
):
    edges = self.agreement_edges(df)
    degree = self.agreement_degree(df)
    return {
        "edges": edges,
        "nodes": degree,
    }


# Most Agreeing Models

def most_agreeing_models(
    self,
    df,
):
    degree = self.agreement_degree(df)
    return degree.sort_values(
        "Degree",
        ascending=False,
    )


# Most Disagreeing Models

def most_disagreeing_models(
    self,
    df,
):
    degree = self.agreement_degree(df)
    return degree.sort_values(
        "Degree",
        ascending=True,
    )


# Model Consensus Score

def model_consensus_score(
    self,
    df,
):
    matrix = self.model_agreement_matrix(df)
    scores = []
    for model in matrix.index:
        value = (
            matrix.loc[model]
            .drop(model)
            .mean()
        )
        scores.append(value)
    return pd.DataFrame(
        {
            "Model": matrix.index,
            "ConsensusScore": np.round(scores, 3),
        }
    ).sort_values(
        "ConsensusScore",
        ascending=False,
    )



# Prediction Diversity

def prediction_diversity(
    self,
    df,
):
    diversity = (
        df
        .groupby(
            [
                "Date",
                "Ticker",
            ]
        )
        ["Prediction"]
        .nunique()
        .reset_index()
    )
    diversity.columns = [
        "Date",
        "Ticker",
        "UniquePredictions",
    ]
    return diversity


# High Disagreement Trades

def high_disagreement_trades(
    self,
    df,
):
    diversity = self.prediction_diversity(df)
    maximum = diversity["UniquePredictions"].max()
    return diversity[
        diversity["UniquePredictions"]
        ==
        maximum
    ]


# Network Summary

def network_summary(
    self,
    df,
):
    edges = self.agreement_edges(df)
    nodes = self.agreement_degree(df)
    return {
        "Nodes":
            len(nodes),
        "Edges":
            len(edges),
        "Average Degree":
            round(
                nodes["Degree"].mean(),
                2,
            ),
        "Maximum Degree":
            int(
                nodes["Degree"].max()
            ),
        "Minimum Degree":
            int(
                nodes["Degree"].min()
            ),
    }


# Consensus Report

def consensus_report(
    self,
    df,
):
    return {
        "Consensus Accuracy":
            self.consensus_accuracy(df),
        "Consensus Distribution":
            self.consensus_distribution(df),
        "Consensus Score":
            self.model_consensus_score(df),
        "Agreement Matrix":
            self.model_agreement_matrix(df),
        "Disagreement Matrix":
            self.model_disagreement_matrix(df),
        "Network":
            self.network_summary(df),
    }


# Prediction Flow Sankey

def prediction_flow_sankey(
    self,
    df,
):
    flow = self.prediction_flow(df)
    labels = sorted(
        set(flow["GroundTruth"]).union(
            set(flow["Prediction"])
        )
    )
    index = {
        label: i
        for i, label in enumerate(labels)
    }
    fig = go.Figure(
        go.Sankey(
            node=dict(
                label=labels,
                pad=20,
                thickness=20,
            ),
            link=dict(
                source=[
                    index[x]
                    for x in flow["GroundTruth"]
                ],
                target=[
                    index[x]
                    for x in flow["Prediction"]
                ],
                value=flow["Count"],
            ),
        )
    )
    fig.update_layout(
        title="Ground Truth → Prediction Flow"
    )
    return fig


# Model Agreement Heatmap

def agreement_heatmap(
    self,
    df,
):
    matrix = self.model_agreement_matrix(df)

    fig = px.imshow(
        matrix,
        text_auto=".2f",
        aspect="auto",
        color_continuous_scale="Viridis",
    )
    fig.update_layout(
        title="Model Agreement Matrix"
    )
    return fig


# Model Disagreement Heatmap

def disagreement_heatmap(
    self,
    df,
):
    matrix = self.model_disagreement_matrix(df)
    fig = px.imshow(
        matrix,
        text_auto=".2f",
        aspect="auto",
        color_continuous_scale="Reds",
    )
    fig.update_layout(
        title="Model Disagreement Matrix"
    )
    return fig


# Consensus Score Bar Chart

def consensus_bar(
    self,
    df,
):
    score = self.model_consensus_score(df)
    fig = px.bar(
        score,
        x="Model",
        y="ConsensusScore",
        color="ConsensusScore",
        text="ConsensusScore",
    )
    fig.update_layout(
        title="Consensus Score"
    )
    return fig


# Agreement Network

def agreement_network(
    self,
    df,
):
    edges = self.agreement_edges(df)
    fig = go.Figure()
    for _, row in edges.iterrows():
        fig.add_trace(
            go.Scatter(
                x=[0, 1],
                y=[0, 1],
                mode="lines+text",
                text=[
                    row["Source"],
                    row["Target"],
                ],
                line=dict(
                    width=6 * row["Weight"]
                ),
                showlegend=False,
            )
        )
    fig.update_layout(
        title="Model Agreement Network"
    )
    return fig

# Prediction Diversity Histogram

def diversity_histogram(
    self,
    df,
):
    diversity = self.prediction_diversity(df)
    fig = px.histogram(
        diversity,
        x="UniquePredictions",
        nbins=10,
    )
    fig.update_layout(
        title="Prediction Diversity"
    )
    return fig


# High Disagreement Table

def show_high_disagreement(
    self,
    df,
):
    st.subheader(
        "High Disagreement Trades"
    )
    st.dataframe(
        self.high_disagreement_trades(df),
        use_container_width=True,
    )


# Consensus Distribution

def consensus_distribution_chart(
    self,
    df,
):
    dist = self.consensus_distribution(df)
    fig = px.pie(
        dist,
        names="Level",
        values="Cases",
    )
    fig.update_layout(
        title="Consensus Distribution"
    )

    return fig


# Model Degree Chart

def agreement_degree_chart(
    self,
    df,
):
    degree = self.agreement_degree(df)
    fig = px.bar(
        degree,
        x="Model",
        y="Degree",
        color="Degree",
        text="Degree",
    )
    fig.update_layout(
        title="Model Connectivity"
    )
    return fig


# Network Statistics Table

def show_network_summary(
    self,
    df,
):
    st.subheader(
        "Consensus Network"
    )
    st.json(
        self.network_summary(df)
    )

# Render Consensus Dashboard

def render_consensus_dashboard(
    self,
    df,
):
    st.header(
        "Consensus Analytics"
    )

    st.plotly_chart(
        self.prediction_flow_sankey(df),
        use_container_width=True,
    )
    st.plotly_chart(
        self.agreement_heatmap(df),
        use_container_width=True,
    )
    st.plotly_chart(
        self.disagreement_heatmap(df),
        use_container_width=True,
    )

    st.plotly_chart(
        self.consensus_bar(df),
        use_container_width=True,
    )

    st.plotly_chart(
        self.diversity_histogram(df),
        use_container_width=True,
    )
    st.plotly_chart(
        self.consensus_distribution_chart(df),
        use_container_width=True,
    )
    st.plotly_chart(
        self.agreement_degree_chart(df),
        use_container_width=True,
    )
    self.show_high_disagreement(df)
    self.show_network_summary(df)



# Executive Metrics

def executive_metrics(
    self,
    df,
):
    """
    Compute executive-level KPIs for the audit report.
    """
    scorecard = self.gt_scorecard(df)
    metrics = {
        "Total Trades":
            len(df),
        "Models":
            df["Model"].nunique(),
        "Assets":
            df["Ticker"].nunique(),
        "Accuracy":
            scorecard["Accuracy"],
        "Average Trust":
            scorecard["Trust"],
        "Average Confidence":
            scorecard["Confidence"],
        "Hallucination Rate":
            scorecard["Hallucination Rate"],
    }
    return metrics


# Best Performing Model

def best_model(
    self,
    df,
):
    leaderboard = self.model_leaderboard(df)
    return leaderboard.iloc[0]


# Worst Performing Model

def worst_model(
    self,
    df,
):
    leaderboard = self.model_leaderboard(df)
    return leaderboard.iloc[-1]


# Best Asset

def best_asset(
    self,
    df,
):
    assets = self.asset_leaderboard(df)
    return assets.iloc[0]


# Worst Asset

def worst_asset(
    self,
    df,
):
    assets = self.asset_leaderboard(df)
    return assets.iloc[-1]


# Top Hallucinating Model

def highest_hallucination_model(
    self,
    df,
):
    ranking = self.hallucination_ranking(df)
    return ranking.iloc[-1]


# Most Trusted Model

def most_trusted_model(
    self,
    df,
):
    trust = self.trust_ranking(df)
    return trust.iloc[0]


# Least Trusted Model

def least_trusted_model(
    self,
    df,
):
    trust = self.trust_ranking(df)
    return trust.iloc[-1]


# Research Findings

def research_findings(
    self,
    df,
):
    findings = []
    best = self.best_model(df)
    findings.append(
        f"Best model: {best['Model']} "
        f"(Score={best['Score']:.3f})"
    )
    worst = self.worst_model(df)
    findings.append(
        f"Weakest model: {worst['Model']} "
        f"(Score={worst['Score']:.3f})"
    )
    trust = self.most_trusted_model(df)
    findings.append(
        f"Highest trust model: "
        f"{trust['Model']} "
        f"({trust['TrustScore']:.3f})"
    )
    hall = self.highest_hallucination_model(df)
    findings.append(
        f"Highest hallucination rate: "
        f"{hall['Model']} "
        f"({hall['HallucinationRate']:.3f})"
    )
    return findings


# Recommendations

def recommendations(
    self,
    df,
):
    rec = []
    score = self.gt_scorecard(df)
    if score["Accuracy"] < 0.70:
        rec.append(
            "Increase model accuracy before deployment."
        )
    if score["Trust"] < 0.70:
        rec.append(
            "Improve calibration and trust estimation."
        )
    if score["Hallucination Rate"] > 0.15:
        rec.append(
            "Reduce hallucination through GT validation."
        )
    if len(rec) == 0:
        rec.append(
            "System is suitable for pilot deployment."
        )
    return rec


# Executive Summary Text

def executive_summary_text(
    self,
    df,
):
    metrics = self.executive_metrics(df)
    best = self.best_model(df)
    summary = f"""

Ground Truth Audit Summary

Total Trades:
{metrics['Total Trades']}

Models Evaluated:
{metrics['Models']}

Assets:
{metrics['Assets']}

Overall Accuracy:
{metrics['Accuracy']:.2%}

Average Trust:
{metrics['Average Trust']:.3f}
Average Confidence:
{metrics['Average Confidence']:.3f}

Best Performing Model:
{best['Model']}
"""
    return summary


# Research Insight Table

def research_insights(
    self,
    df,
):
    findings = self.research_findings(df)
    recommendations = self.recommendations(df)
    rows = []
    for f in findings:
        rows.append(
            {
                "Category":
                    "Finding",
                "Description":
                    f,
            }
        )
    for r in recommendations:
        rows.append(
            {
                "Category":
                    "Recommendation",
                "Description":
                    r,
            }
        )
    return pd.DataFrame(rows)


# Audit KPI Table

def audit_kpi_table(
    self,
    df,
):
    metrics = self.executive_metrics(df)
    return pd.DataFrame(
        {
            "Metric":
                list(metrics.keys()),
            "Value":
                list(metrics.values()),
        }
    )


# Show Executive Report

def show_executive_report(
    self,
    df,
):
    st.header(
        "Executive Audit Summary"
    )
    st.text(
        self.executive_summary_text(df)
    )
    st.subheader(
        "Key Performance Indicators"
    )
    st.dataframe(
        self.audit_kpi_table(df),
        use_container_width=True,
    )
    st.subheader(
        "Research Findings"
    )
    st.dataframe(
        self.research_insights(df),
        use_container_width=True,
    )

# JSON Summary

def executive_json(
    self,
    df,
):
    return {
        "metrics":
            self.executive_metrics(df),
        "findings":
            self.research_findings(df),
        "recommendations":
            self.recommendations(df),
    }

from datetime import datetime
import html

# HTML CSS Theme

def html_theme(self):

    return """
    <style>
    body{
        font-family:Arial,Helvetica,sans-serif;
        margin:40px;
        background:#ffffff;
        color:#222;
    }
    h1{
        color:#0B5394;
        border-bottom:3px solid #0B5394;
        padding-bottom:8px;
    }
    h2{
        color:#1F4E79;
        margin-top:35px;
    }
    table{
        border-collapse:collapse;
        width:100%;
        margin-top:10px;
        margin-bottom:20px;
    }
    table,th,td{
        border:1px solid #d0d0d0;
    }
    th{
        background:#EAF2F8;
        padding:8px;
        text-align:left;
    }
    td{
        padding:7px;
    }
    .metric{
        display:inline-block;
        width:220px;
        margin:10px;
        padding:15px;
        border-radius:6px;
        background:#F5F5F5;
        border-left:6px solid #0B5394;
    }
    .footer{
        margin-top:60px;
        color:gray;
        font-size:12px;
    }
    </style>
    """


# DataFrame -> HTML

def dataframe_to_html(
    self,
    df,
):
    return df.to_html(
        index=False,
        border=0,
        classes="table",
        justify="left",
    )


# KPI Cards HTML

def metrics_html(
    self,
    df,
):
    metrics = self.executive_metrics(df)
    html_text = ""
    for key, value in metrics.items():
        html_text += f"""
        <div class="metric">
        <b>{html.escape(str(key))}</b><br>
        <h2>{html.escape(str(value))}</h2>
        </div>
        """
    return html_text


# Findings HTML

def findings_html(
    self,
    df,
):
    findings = self.research_findings(df)
    text = "<ul>"
    for item in findings:
        text += f"<li>{html.escape(item)}</li>"
    text += "</ul>"
    return text


# Recommendation HTML

def recommendations_html(
    self,
    df,
):
    rec = self.recommendations(df)
    text = "<ul>"
    for item in rec:
        text += f"<li>{html.escape(item)}</li>"
    text += "</ul>"
    return text


# Publication Tables

def publication_tables(
    self,
    df,
):
    return {
        "Executive KPIs":
            self.audit_kpi_table(df),
        "Model Leaderboard":
            self.model_leaderboard(df),
        "Asset Leaderboard":
            self.asset_leaderboard(df),
        "Hallucination Ranking":
            self.hallucination_ranking(df),
        "Calibration Ranking":
            self.calibration_ranking(df),
    }


# Publication HTML

def publication_tables_html(
    self,
    df,
):
    tables = self.publication_tables(df)
    text = ""
    for name, table in tables.items():
        text += f"<h2>{html.escape(name)}</h2>"
        text += self.dataframe_to_html(table)
    return text


# Markdown Report

def markdown_report(
    self,
    df,
):
    metrics = self.executive_metrics(df)
    report = "# Ground Truth Audit Report\n\n"
    report += f"Generated: {datetime.now()}\n\n"
    report += "## Executive Metrics\n\n"
    for k, v in metrics.items():
        report += f"- **{k}:** {v}\n"
    report += "\n"
    report += "## Research Findings\n\n"
    for item in self.research_findings(df):
        report += f"- {item}\n"
    report += "\n"
    report += "## Recommendations\n\n"
    for item in self.recommendations(df):
        report += f"- {item}\n"
    return report


# HTML Report Body

def html_report_body(
    self,
    df,
):
    body = f"""
    <h1>
    Ground Truth Audit Report
    </h1>
    <p>
    Generated:
    {datetime.now()}
    </p>
    <h2>
    Executive Metrics
    </h2>
    {self.metrics_html(df)}
    <h2>
    Research Findings
    </h2>
    {self.findings_html(df)}
    <h2>
    Recommendations
    </h2>
    {self.recommendations_html(df)}
    <h2>
    Publication Tables
    </h2>
    {self.publication_tables_html(df)}
    """
    return body


# Complete HTML String

def html_report(
    self,
    df,
):
    html_doc = f"""
    <html>
    <head>
    {self.html_theme()}
    </head>
    <body>
    {self.html_report_body(df)}
    <div class="footer">
    Generated by GTTableVisualizer
    </div>
    </body>
    </html>
    """
    return html_doc


# Preview HTML

def preview_html_report(
    self,
    df,
):
    st.components.v1.html(
        self.html_report(df),
        height=900,
        scrolling=True,
    )


# Save HTML Report

def save_html_report(
    self,
    df,
    filepath="audit_report.html",
):
    """
    Save complete HTML report to disk.
    Returns
    -------
    str
        Saved file path.
    """
    report = self.html_report(df)
    with open(
        filepath,
        "w",
        encoding="utf-8",
    ) as f:
        f.write(report)
    return filepath


# Plotly Figure -> HTML

def plotly_to_html(
    self,
    figure,
):
    """
    Convert a Plotly figure into an
    embeddable HTML fragment.
    """
    if figure is None:
        return ""
    return figure.to_html(
        include_plotlyjs=False,
        full_html=False,
    )


# Embed Plotly Figure

def embed_plotly_figure(
    self,
    title,
    figure,
):
    """
    Wrap Plotly figure inside
    a publication-ready HTML section.
    """
    return f"""
    <section
        style="margin-top:35px;">
        <h2>{html.escape(title)}</h2>
        {self.plotly_to_html(figure)}
    </section>
    """

# Generic Report Section

def report_section(
    self,
    title,
    content,
):
    """
    Generic HTML section.
    """
    return f"""
    <section
        style="margin-top:40px;">
        <h2>
        {html.escape(title)}
        </h2>
        {content}
    </section>
    """

# HTML Table Section

def html_table_section(
    self,
    title,
    dataframe,
):
    """
    Convert dataframe into
    a styled HTML section.
    """
    table = self.dataframe_to_html(
        dataframe
    )
    return self.report_section(
        title,
        table,
    )


# HTML Text Section

def html_text_section(
    self,
    title,
    text,
):
    """
    Plain text HTML section.
    """
    return self.report_section(
        title,
        f"""
        <p>
        {html.escape(text)}
        </p>
        """,
    )

# HTML Bullet List

def html_bullet_list(
    self,
    title,
    items,
):
    """
    Create unordered list section.
    """
    html_list = "<ul>"
    for item in items:
        html_list += (
            "<li>"
            f"{html.escape(str(item))}"
            "</li>"
        )
    html_list += "</ul>"
    return self.report_section(
        title,
        html_list,
    )


# Report Header

def report_header(
    self,
    title="Ground Truth Audit Report",
):
    """
    HTML report header.
    """
    return f"""
    <header>
        <h1>
        {html.escape(title)}
        </h1>
        <p>
        Generated:
        {datetime.now():%Y-%m-%d %H:%M:%S}
        </p>
        <hr>
    </header>
    """

# Build HTML Sections

def build_html_sections(
    self,
    df,
):
    """
    Build all major HTML sections used in the
    publication report.
    """
    sections = []

    # Executive Summary
    sections.append(
        self.html_text_section(
            "Executive Summary",
            self.executive_summary_text(df),
        )
    )

    # KPI Table
    sections.append(
        self.html_table_section(
            "Key Performance Indicators",
            self.audit_kpi_table(df),
        )
    )

    # Research Insights
    sections.append(
        self.html_table_section(
            "Research Insights",
            self.research_insights(df),
        )
    )

    # Leaderboard
    sections.append(
        self.html_table_section(
            "Model Leaderboard",
            self.model_leaderboard(df),
        )
    )

    # Asset Leaderboard
    sections.append(
        self.html_table_section(
            "Asset Leaderboard",
            self.asset_leaderboard(df),
        )
    )

    # Findings
    sections.append(
        self.html_bullet_list(
            "Research Findings",
            self.research_findings(df),
        )
    )

    # Recommendations
    sections.append(
        self.html_bullet_list(
            "Recommendations",
            self.recommendations(df),
        )
    )
    return sections


# Assemble HTML Report

def assemble_html_report(
    self,
    df,
):
    """
    Assemble a complete HTML report.
    """
    report = []
    report.append("<html>")
    report.append("<head>")
    report.append(self.html_theme())
    report.append("</head>")
    report.append("<body>")
    report.append(
        self.report_header()
    )
    for section in self.build_html_sections(df):
        report.append(section)
    report.append(
        """
        <footer
        style="margin-top:60px;
               font-size:12px;
               color:gray;
               border-top:1px solid #ccc;
               padding-top:20px;">

        Generated automatically by
        GTTableVisualizer.

        </footer>
        """
    )
    report.append("</body>")
    report.append("</html>")
    return "\n".join(report)

# Export HTML

def export_html(
    self,
    df,
    filepath="audit_report.html",
):
    html_text = self.assemble_html_report(df)
    with open(
        filepath,
        "w",
        encoding="utf-8",
    ) as f:
        f.write(html_text)
    return filepath


# HTML Download Button

def html_download_button(
    self,
    df,
):
    html_text = self.assemble_html_report(df)
    st.download_button(
        label="📄 Download HTML Report",
        data=html_text,
        file_name="audit_report.html",
        mime="text/html",
    )

# Preview HTML Report

def preview_full_report(
    self,
    df,
):
    st.components.v1.html(
        self.assemble_html_report(df),
        height=900,
        scrolling=True,
    )


# Report Dashboard

def render_report_dashboard(
    self,
    df,
):
    st.title("Publication Report")
    st.markdown(
        """
        Generate a publication-ready audit report
        for Ground Truth (GT) evaluation.
        """
    )
    self.html_download_button(df)
    st.divider()
    self.preview_full_report(df)


# Initialize PDF Styles

def init_pdf_styles(self):
    """
    Create reusable ReportLab paragraph styles.
    """
    styles = getSampleStyleSheet()
    custom = {
        "title": ParagraphStyle(
            "GTTitle",
            parent=styles["Title"],
            alignment=TA_CENTER,
            fontSize=24,
            leading=30,
            textColor=colors.darkblue,
            spaceAfter=24,
        ),
        "heading": ParagraphStyle(
            "GTHeading",
            parent=styles["Heading2"],
            alignment=TA_LEFT,
            fontSize=16,
            leading=20,
            textColor=colors.darkblue,
            spaceBefore=16,
            spaceAfter=10,
        ),
        "normal": ParagraphStyle(
            "GTNormal",
            parent=styles["BodyText"],
            fontSize=11,
            leading=16,
            spaceAfter=8,
        ),
        "small": ParagraphStyle(
            "GTSmall",
            parent=styles["BodyText"],
            fontSize=9,
            leading=12,
            textColor=colors.grey,
        ),
    }
    return custom


# PDF Metadata

def pdf_metadata(self):
    """
    Metadata dictionary for report.
    """
    return {
        "title":
            "Ground Truth Audit Report",
        "author":
            "GTTableVisualizer",
        "subject":
            "LLM Trading Audit",
        "keywords":
            "LLM,Trading,Audit,Ground Truth,Explainability",
        "created":
            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            ),
    }

# PDF Cover Page

def pdf_cover_page(self):
    """
    Returns ReportLab flowables for
    title page.
    """
    styles = self.init_pdf_styles()
    meta = self.pdf_metadata()
    story = []
    story.append(
        Spacer(
            1,
            1.2 * inch,
        )
    )
    story.append(
        Paragraph(
            meta["title"],
            styles["title"],
        )
    )
    story.append(
        Paragraph(
            "<b>Auditing LLM Trading:</b><br/>"
            "Bridging Theory and Market Reality "
            "with the GT Table",
            styles["heading"],
        )
    )

    story.append(
        Spacer(
            1,
            0.4 * inch,
        )
    )

    story.append(
        Paragraph(
            f"<b>Generated:</b> "
            f"{meta['created']}",
            styles["normal"],
        )
    )

    story.append(
        Paragraph(
            f"<b>Author:</b> "
            f"{meta['author']}",
            styles["normal"],
        )
    )
    story.append(
        Paragraph(
            f"<b>Subject:</b> "
            f"{meta['subject']}",
            styles["normal"],
        )
    )
    story.append(
        Spacer(
            1,
            0.6 * inch,
        )
    )

    story.append(

        Paragraph(

            """
            This report provides a comprehensive
            audit of Large Language Model (LLM)
            trading systems using Ground Truth
            (GT) validation, trust analysis,
            calibration assessment, model
            comparison, hallucination analysis,
            and explainability metrics.
            """,
            styles["normal"],
        )
    )
    story.append(
        Spacer(
            1,
            0.8 * inch,
        )
    )
    story.append(
        Paragraph(
            "Prepared automatically by "
            "GTTableVisualizer",
            styles["small"],
        )
    )
    story.append(
        PageBreak()
    )
    return story


# PDF Executive Summary Page

def pdf_executive_summary(self, df):
    """
    Build executive summary page.
    """
    styles = self.init_pdf_styles()
    story = []
    story.append(
        Paragraph(
            "Executive Summary",
            styles["heading"],
        )
    )
    story.append(
        Paragraph(
            self.executive_summary_text(df),
            styles["normal"],
        )
    )
    story.append(
        Spacer(
            1,
            0.2 * inch,
        )
    )
    return story


# PDF KPI Table

def pdf_kpi_table(self, df):
    """
    Build KPI table for PDF.
    """
    styles = self.init_pdf_styles()
    kpi = self.audit_kpi_table(df)
    table_data = [list(kpi.columns)]
    table_data.extend(
        kpi.values.tolist()
    )
    table = Table(
        table_data,
        hAlign="LEFT",
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND",
                 (0, 0),
                 (-1, 0),
                 colors.HexColor("#1F4E79")),
                ("TEXTCOLOR",
                 (0, 0),
                 (-1, 0),
                 colors.white),
                ("GRID",
                 (0, 0),
                 (-1, -1),
                 0.5,
                 colors.grey),
                ("FONTNAME",
                 (0, 0),
                 (-1, 0),
                 "Helvetica-Bold"),
                ("BACKGROUND",
                 (0, 1),
                 (-1, -1),
                 colors.whitesmoke),
                ("BOTTOMPADDING",
                 (0, 0),
                 (-1, 0),
                 10),
            ]
        )
    )
    story = [
        Paragraph(
            "Key Performance Indicators",
            styles["heading"],
        ),
        table,
        Spacer(1, 0.25 * inch),
    ]
    return story


# PDF Findings

def pdf_findings(self, df):
    """
    Research findings section.
    """
    styles = self.init_pdf_styles()
    story = [
        Paragraph(
            "Research Findings",
            styles["heading"],
        )
    ]
    for item in self.research_findings(df):
        story.append(
            Paragraph(
                f"• {item}",
                styles["normal"],
            )
        )
    story.append(
        Spacer(1, 0.2 * inch)
    )
    return story

# PDF Recommendations

def pdf_recommendations(self, df):
    """
    Recommendations section.
    """
    styles = self.init_pdf_styles()
    story = [
        Paragraph(
            "Recommendations",
            styles["heading"],
        )
    ]
    for item in self.recommendations(df):
        story.append(
            Paragraph(
                f"• {item}",
                styles["normal"],
            )
        )
    story.append(
        Spacer(1, 0.25 * inch)
    )
    return story


# Build PDF Report

def build_pdf_report(
    self,
    df,
    filepath="audit_report.pdf",
):
    """
    Generate complete PDF report.
    """
    doc = SimpleDocTemplate(
        filepath,
        pagesize=letter,
    )
    story = []
    story.extend(
        self.pdf_cover_page()
    )
    story.extend(
        self.pdf_executive_summary(df)
    )
    story.extend(

        self.pdf_kpi_table(df)
    )
    story.extend(
        self.pdf_findings(df)
    )
    story.extend(
        self.pdf_recommendations(df)
    )
    doc.build(story)
    return filepath


# Streamlit PDF Download

def pdf_download_button(
    self,
    df,
):
    """
    Create PDF in a temporary file and expose
    it through a Streamlit download button.
    """

    import tempfile
    with tempfile.NamedTemporaryFile(
        suffix=".pdf",
        delete=False,
    ) as tmp:
        path = self.build_pdf_report(
            df,
            filepath=tmp.name,
        )
    with open(path, "rb") as f:
        pdf_bytes = f.read()
    st.download_button(
        label="📄 Download PDF Report",
        data=pdf_bytes,
        file_name="audit_report.pdf",
        mime="application/pdf",
    )


# Create Excel Workbook

def create_excel_workbook(self):
    """
    Create a new OpenPyXL workbook.
    """
    workbook = Workbook()
    workbook.properties.creator = "GTTableVisualizer"
    workbook.properties.title = "Ground Truth Audit Report"
    workbook.properties.subject = "LLM Trading Audit"
    workbook.properties.description = (
        "Automatically generated audit workbook."
    )
    return workbook


# Excel Styles

def excel_styles(self):
    """
    Common styles used throughout workbook.
    """
    styles = {
        "header_font": Font(
            bold=True,
            color="FFFFFF",
            size=12,
        ),
        "header_fill": PatternFill(
            fill_type="solid",
            fgColor="1F4E79",
        ),
        "body_font": Font(
            size=11,
        ),
        "center": Alignment(
            horizontal="center",
            vertical="center",
        ),
        "thin_border": Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        ),
        "title_font": Font(
            bold=True,
            size=16,
        ),
    }
    return styles

# Auto Fit Columns

def autofit_columns(
    self,
    worksheet,
):
    """
    Automatically adjust column widths.
    """
    for column in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(
            column[0].column
        )
        for cell in column:
            try:
                value = str(cell.value)
                if len(value) > max_length:
                    max_length = len(value)
            except Exception:
                pass
        worksheet.column_dimensions[
            column_letter
        ].width = max_length + 3


# Apply Header Style

def style_header(
    self,
    worksheet,
):
    styles = self.excel_styles()
    for cell in worksheet[1]:
        cell.font = styles["header_font"]
        cell.fill = styles["header_fill"]
        cell.alignment = styles["center"]
        cell.border = styles["thin_border"]


# DataFrame -> Worksheet

def dataframe_to_worksheet(
    self,
    worksheet,
    dataframe,
):
    """
    Write DataFrame into worksheet.
    """
    # Header
    worksheet.append(
        dataframe.columns.tolist()
    )
    # Data
    for row in dataframe.itertuples(index=False):
        worksheet.append(
            list(row)
        )

    # Style

    self.style_header(
        worksheet
    )
    styles = self.excel_styles()
    for row in worksheet.iter_rows(
        min_row=2
    ):
        for cell in row:
            cell.font = styles["body_font"]
            cell.border = styles["thin_border"]
    self.autofit_columns(
        worksheet
    )
    worksheet.freeze_panes = "A2"

# Create Worksheet

def add_dataframe_sheet(
    self,
    workbook,
    sheet_name,
    dataframe,
):
    """
    Add a new worksheet from a DataFrame.
    """
    worksheet = workbook.create_sheet(
        title=sheet_name
    )
    self.dataframe_to_worksheet(
        worksheet,
        dataframe,
    )
    return worksheet

# Remove Default Sheet

def remove_default_sheet(
    self,
    workbook,
):
    if "Sheet" in workbook.sheetnames:
        workbook.remove(
            workbook["Sheet"]
        )
    return workbook


# Executive Summary Sheet

def add_executive_summary_sheet(
    self,
    workbook,
    df,
):
    """
    Create Executive Summary worksheet.
    """
    ws = workbook.create_sheet(
        title="Executive Summary"
    )
    styles = self.excel_styles()
    ws["A1"] = "Ground Truth Audit Report"
    ws["A1"].font = styles["title_font"]
    ws["A3"] = "Generated"
    ws["B3"] = datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )
    metrics = self.executive_metrics(df)
    start_row = 6
    ws.cell(
        row=start_row,
        column=1,
        value="Metric",
    )
    ws.cell(
        row=start_row,
        column=2,
        value="Value",
    )
    self.style_header(ws)
    row = start_row + 1
    for key, value in metrics.items():
        ws.cell(
            row=row,
            column=1,
            value=key,
        )
        ws.cell(
            row=row,
            column=2,
            value=value,
        )
        row += 1
    ws.cell(
        row=row + 2,
        column=1,
        value="Best Model",
    )
    ws.cell(
        row=row + 2,
        column=2,
        value=self.best_model(df)["Model"],
    )
    ws.cell(
        row=row + 3,
        column=1,
        value="Most Trusted Model",
    )
    ws.cell(
        row=row + 3,
        column=2,
        value=self.most_trusted_model(df)["Model"],
    )
    ws.freeze_panes = "A6"
    self.autofit_columns(ws)
    return ws


# KPI Sheet

def add_kpi_sheet(
    self,
    workbook,
    df,
):
    """
    Create KPI worksheet.
    """
    worksheet = workbook.create_sheet(
        title="KPI Dashboard"
    )
    kpi = self.audit_kpi_table(df)
    self.dataframe_to_worksheet(
        worksheet,
        kpi,
    )
    worksheet.sheet_view.showGridLines = True
    worksheet.freeze_panes = "A2"
    return worksheet


# Model Leaderboard Sheet

def add_model_leaderboard_sheet(
    self,
    workbook,
    df,
):
    """
    Create Model Leaderboard worksheet.
    """
    worksheet = workbook.create_sheet(
        title="Model Leaderboard"
    )

    leaderboard = self.model_leaderboard(df)

    self.dataframe_to_worksheet(
        worksheet,
        leaderboard,
    )
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )
    return worksheet


# Workbook Summary

def add_workbook_summary(
    self,
    workbook,
):
    """
    Create workbook information sheet.
    """
    ws = workbook.create_sheet(
        title="Workbook Info"
    )

    styles = self.excel_styles()

    ws["A1"] = "Workbook Information"
    ws["A1"].font = styles["title_font"]
    rows = [
        (
            "Project",
            "Auditing LLM Trading",
        ),
        (
            "Component",
            "GTTableVisualizer",
        ),
        (
            "Generated",
            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            ),
        ),
        (
            "Workbook Version",
            "1.0",
        ),
    ]
    r = 3
    for key, value in rows:
        ws.cell(
            row=r,
            column=1,
            value=key,
        )
        ws.cell(
            row=r,
            column=2,
            value=value,
        )
        r += 1
    self.autofit_columns(ws)
    return ws


# Asset Leaderboard Sheet

def add_asset_leaderboard_sheet(
    self,
    workbook,
    df,
):
    """
    Asset performance worksheet.
    """
    worksheet = workbook.create_sheet(
        title="Asset Leaderboard"
    )
    assets = self.asset_leaderboard(df)
    self.dataframe_to_worksheet(
        worksheet,
        assets,
    )
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    return worksheet


# Hallucination Analysis Sheet

def add_hallucination_sheet(
    self,
    workbook,
    df,
):
    """
    Hallucination ranking worksheet.
    """
    worksheet = workbook.create_sheet(
        title="Hallucination"
    )
    hallucination = self.hallucination_ranking(df)
    self.dataframe_to_worksheet(
        worksheet,
        hallucination,
    )
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    return worksheet


# Calibration Sheet

def add_calibration_sheet(
    self,
    workbook,
    df,
):
    """
    Calibration ranking worksheet.
    """
    worksheet = workbook.create_sheet(
        title="Calibration"
    )
    calibration = self.calibration_ranking(df)
    self.dataframe_to_worksheet(
        worksheet,
        calibration,
    )
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    return worksheet


# Consensus Sheet

def add_consensus_sheet(
    self,
    workbook,
    df,
):
    """
    Consensus metrics worksheet.
    """
    worksheet = workbook.create_sheet(
        title="Consensus"
    )
    consensus = self.model_consensus_score(df)
    self.dataframe_to_worksheet(
        worksheet,
        consensus,
    )
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    return worksheet

# Final Formatting

def finalize_excel_workbook(
    self,
    workbook,
):
    """
    Apply workbook-wide formatting.
    """
    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = True
        if sheet.max_row > 1:
            sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        self.autofit_columns(sheet)
    return workbook


# Build Complete Workbook

def build_excel_report(
    self,
    df,
):
    """
    Build complete Excel workbook.
    """
    workbook = self.create_excel_workbook()
    self.remove_default_sheet(workbook)
    self.add_workbook_summary(
        workbook
    )
    self.add_executive_summary_sheet(
        workbook,
        df,
    )
    self.add_kpi_sheet(
        workbook,
        df,
    )
    self.add_model_leaderboard_sheet(
        workbook,
        df,
    )
    self.add_asset_leaderboard_sheet(
        workbook,
        df,
    )
    self.add_hallucination_sheet(
        workbook,
        df,
    )
    self.add_calibration_sheet(
        workbook,
        df,
    )
    self.add_consensus_sheet(
        workbook,
        df,
    )
    self.finalize_excel_workbook(
        workbook
    )
    return workbook


# Save Workbook

def save_excel_report(
    self,
    df,
    filepath="audit_report.xlsx",
):
    """
    Save workbook to disk.
    """
    workbook = self.build_excel_report(df)
    workbook.save(filepath)
    return filepath


# Workbook Summary

def workbook_summary(
    self,
    workbook,
):
    """
    Return workbook metadata.
    """
    return {
        "Sheets": len(workbook.sheetnames),
        "Sheet Names": workbook.sheetnames,
        "Creator":
            workbook.properties.creator,
        "Title":
            workbook.properties.title,
        "Subject":
            workbook.properties.subject,
    }


# Validate Workbook

def validate_workbook(
    self,
    workbook,
):
    """
    Validate workbook before saving.
    """
    errors = []
    if len(workbook.sheetnames) == 0:
        errors.append(
            "Workbook has no sheets."
        )
    for sheet in workbook.worksheets:
        if sheet.max_row == 0:
            errors.append(
                f"{sheet.title} is empty."
            )
    return {
        "valid": len(errors) == 0,
        "errors": errors,
    }


# Export Workbook

def export_excel_report(
    self,
    df,
    filepath="audit_report.xlsx",
):
    """
    Build and export workbook.
    """
    workbook = self.build_excel_report(df)
    validation = self.validate_workbook(
        workbook
    )
    if not validation["valid"]:
        raise ValueError(
            "\n".join(
                validation["errors"]
            )
        )
    workbook.save(filepath)
    return filepath

# Streamlit Download Button

def excel_download_button(
    self,
    df,
):
    """
    Download Excel report.
    """

    import tempfile
    import os
    with tempfile.TemporaryDirectory() as tmp:
        filename = os.path.join(
            tmp,
            "audit_report.xlsx",
        )
        self.export_excel_report(
            df,
            filename,
        )
        with open(
            filename,
            "rb",
        ) as f:
            data = f.read()
    st.download_button(
        label="📊 Download Excel Report",
        data=data,
        file_name="audit_report.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# Show Workbook Summary

def show_workbook_summary(
    self,
    df,
):
    """
    Display workbook information.
    """
    workbook = self.build_excel_report(df)
    summary = self.workbook_summary(
        workbook
    )
    st.subheader(
        "Workbook Summary"
    )
    st.json(summary)


# Render Excel Export Page

def render_excel_export(
    self,
    df,
):
    """
    Streamlit Excel export page.
    """
    st.header(
        "Excel Report Export"
    )
    st.markdown(
        """
        Export the complete Ground Truth
        Audit Workbook.
        """
    )
    self.show_workbook_summary(
        df
    )
    st.divider()
    self.excel_download_button(
        df
    )


# Export Manager

def export_manager(
    self,
    df,
):
    """
    Central export interface.
    """
    st.title(
        "Report Export Center"
    )
    tab1, tab2 = st.tabs(
        [
            "PDF",
            "Excel",
        ]
    )
    with tab1:
        self.pdf_download_button(
            df
        )
    with tab2:
        self.render_excel_export(
            df
        )


# Report Information

def report_information(
    self,
):
    """
    Static report metadata.
    """
    return {
        "Project":
            "Auditing LLM Trading",
        "Component":
            "GTTableVisualizer",
        "Version":
            "1.0.0",
        "Exports":
            [
                "HTML",
                "PDF",
                "Excel",
            ],
    }


# Export DataFrame as CSV

def export_csv(
    self,
    dataframe,
    filepath,
):
    """
    Export DataFrame to CSV.
    """
    dataframe.to_csv(
        filepath,
        index=False,
    )
    return filepath


# Export All CSV Files

def export_csv_package(
    self,
    df,
    output_dir,
):
    """
    Export all audit tables as CSV.
    """
    output_dir = Path(output_dir)
    output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )
    files = {}
    tables = {
        "audit_kpis.csv":
            self.audit_kpi_table(df),
        "model_leaderboard.csv":
            self.model_leaderboard(df),
        "asset_leaderboard.csv":
            self.asset_leaderboard(df),
        "hallucination.csv":
            self.hallucination_ranking(df),
        "calibration.csv":
            self.calibration_ranking(df),
        "consensus.csv":
            self.model_consensus_score(df),
        "research_insights.csv":
            self.research_insights(df),
    }
    for filename, table in tables.items():
        path = output_dir / filename
        self.export_csv(
            table,
            path,
        )
        files[filename] = str(path)
    return files


# Export Manifest

def export_manifest(
    self,
):
    """
    Metadata describing exported package.
    """
    return {
        "project":
            "Auditing LLM Trading",
        "component":
            "GTTableVisualizer",
        "version":
            "1.0.0",
        "exports":
            [
                "CSV",
                "HTML",
                "PDF",
                "Excel",
            ],
    }


# Save Manifest

def save_manifest(
    self,
    output_dir,
):
    """
    Save manifest.json
    """
    path = Path(output_dir) / "manifest.json"
    with open(
        path,
        "w",
        encoding="utf-8",
    ) as f:
        json.dump(
            self.export_manifest(),
            f,
            indent=4,
        )
    return path


# Create ZIP Package

def create_zip_package(
    self,
    folder,
    zip_path,
):
    """
    Compress folder into ZIP.
    """
    folder = Path(folder)
    with zipfile.ZipFile(
        zip_path,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as archive:
        for file in folder.rglob("*"):
            if file.is_file():
                archive.write(
                    file,
                    arcname=file.relative_to(folder),
                )
    return zip_path


# Build Complete Export Package

def build_export_package(
    self,
    df,
    output_dir,
):
    """
    Generate all reports.
    """
    output_dir = Path(output_dir)
    output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )
    # HTML
    self.export_html(
        df,
        output_dir / "audit_report.html",
    )
    # PDF
    self.build_pdf_report(
        df,
        output_dir / "audit_report.pdf",
    )
    # Excel
    self.export_excel_report(
        df,
        output_dir / "audit_report.xlsx",
    )

    # CSV
    self.export_csv_package(
        df,
        output_dir,
    )
    # Manifest
    self.save_manifest(
        output_dir,
    )
    return output_dir


# Export ZIP Package

def export_zip_package(
    self,
    df,
    filepath="audit_package.zip",
):
    """
    Create ZIP package.
    """
    with tempfile.TemporaryDirectory() as tmp:
        self.build_export_package(
            df,
            tmp,
        )
        self.create_zip_package(
            tmp,
            filepath,
        )
    return filepath


# Package Summary

def package_summary(
    self,
):
    """
    Return package contents.
    """
    return {
        "Reports":
            [
                "HTML",
                "PDF",
                "Excel",
            ],
        "Tables":
            [
                "KPIs",
                "Leaderboards",
                "Consensus",
                "Calibration",
                "Hallucination",
            ],
        "CSV Files":
            7,
    }


# Validate Export Package

def validate_export_package(
    self,
    package_path,
):
    """
    Validate that the generated ZIP package exists
    and is non-empty.
    """
    if not os.path.exists(package_path):
        return {
            "valid": False,
            "message": "Package does not exist."
        }
    if os.path.getsize(package_path) == 0:
        return {
            "valid": False,
            "message": "Package is empty."
        }
    return {
        "valid": True,
        "message": "Package validation successful."
    }


# Build and Read ZIP Package

def build_zip_bytes(
    self,
    df,
):
    """
    Build ZIP package and return bytes.
    """
    with tempfile.TemporaryDirectory() as tmpdir:
        zip_path = os.path.join(
            tmpdir,
            "GT_Audit_Report.zip",
        )
        self.export_zip_package(
            df,
            filepath=zip_path,
        )
        validation = self.validate_export_package(
            zip_path
        )
        if not validation["valid"]:
            raise RuntimeError(
                validation["message"]
            )
        with open(
            zip_path,
            "rb",
        ) as f:
            return f.read()


# Download ZIP Button

def zip_download_button(
    self,
    df,
):
    """
    Streamlit ZIP download button.
    """
    zip_bytes = self.build_zip_bytes(df)
    st.download_button(
        label="📦 Download Complete Audit Package",
        data=zip_bytes,
        file_name="GT_Audit_Report.zip",
        mime="application/zip",
    )


# Export Status Panel

def export_status_panel(
    self,
):
    """
    Display export status.
    """
    st.success(
        "All export modules are available."
    )
    st.markdown(
        """
        **Supported formats**
        - HTML Report
        - PDF Report
        - Excel Workbook
        - CSV Tables
        - ZIP Package
        """
    )


# Export Information

def export_information(
    self,
):
    """
    Display export metadata.
    """
    info = self.report_information()
    st.subheader(
        "Export Information"
    )
    st.json(info)


# Export Progress

def export_progress(
    self,
):
    """
    Simulated export progress.
    """
    progress = st.progress(0)
    for value in range(0, 101, 20):
        progress.progress(value)
    progress.empty()


# Unified Export Dashboard

def render_export_dashboard(
    self,
    df,
):
    """
    Render export page.
    """
    st.title(
        "Export Center"
    )
    self.export_information()
    st.divider()
    self.export_status_panel()
    st.divider()
    self.export_progress()
    col1, col2 = st.columns(2)
    with col1:
        self.pdf_download_button(df)
        self.html_download_button(df)
    with col2:
        self.excel_download_button(df)
        self.zip_download_button(df)


# Export All Reports

def export_all_reports(
    self,
    df,
):
    """
    Main export entry point.
    """
    self.render_export_dashboard(df)
